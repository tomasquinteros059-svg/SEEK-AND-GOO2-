"""Genera la plantilla de carga de entregas y el archivo de prueba de 50 filas.

Uso:  python3 herramientas/generar-planillas.py
Crea:
  apps/web/public/plantilla-entregas.xlsx   la que descarga el administrador (con 2 filas de ejemplo)
  apps/web/e2e/archivos/entregas-50.xlsx    la de la prueba de aceptación de la Fase 2 (41 listas, 5 por revisar, 4 con error)

Todos los clientes, direcciones, teléfonos y guías son ficticios.
"""
from datetime import date, time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ENCABEZADOS = ["Cliente (local)", "Dirección", "Comuna", "Contacto recepción", "Teléfono contacto", "Fecha",
               "Hora comprometida", "Recepción desde", "Recepción hasta", "N° guía", "Orden de compra", "Bultos",
               "Cobrar al entregar", "Monto", "Forma de pago", "Conductor", "Notas"]


def libro(filas, instrucciones=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Entregas"
    ws.append(ENCABEZADOS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="143B5C")
    for f in filas:
        ws.append(f)
    for fila in ws.iter_rows(min_row=2):
        fila[5].number_format = "DD-MM-YYYY"
        for i in (6, 7, 8):
            fila[i].number_format = "HH:MM"
    for col, ancho in zip("ABCDEFGHIJKLMNOPQ", [24, 30, 14, 20, 16, 12, 10, 10, 10, 12, 14, 8, 10, 10, 14, 18, 30]):
        ws.column_dimensions[col].width = ancho
    dv = DataValidation(type="list", formula1='"No,Sí"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("M2:M1000")
    ws.freeze_panes = "A2"
    if instrucciones:
        i = wb.create_sheet("Instrucciones")
        for linea in [
            "Cómo llenar esta planilla",
            "",
            "Obligatorio: Cliente (local), Dirección, Fecha, N° guía y Bultos.",
            "Cliente (local): el nombre del local que recibe, por ejemplo «Jumbo La Florida». Si ya existe, se reconoce solo.",
            "Comuna: una de la Región Metropolitana. Define el sector.",
            "Teléfono contacto: el de quien recibe en el local. Sin él no se le podrá avisar por WhatsApp.",
            "Fecha: 27-10-2026. Horas: 09:30.",
            "Recepción desde / hasta: el horario en que el local recibe mercadería.",
            "N° guía: la guía de despacho. No se puede repetir mientras la entrega esté vigente.",
            "Cobrar al entregar: Sí o No. Si es Sí, escribe el Monto y la Forma de pago (transferencia o efectivo).",
            "Conductor: su nombre o su correo, tal como está en Equipo. Si lo dejas vacío, la asignas después.",
            "",
            "Antes de crear nada verás una vista previa: listas, por revisar y con error.",
        ]:
            i.append([linea])
        i["A1"].font = Font(bold=True, size=14)
        i.column_dimensions["A"].width = 110
    return wb


ejemplo = [
    ["Jumbo La Florida", "Av. Vicuña Mackenna 7110", "La Florida", "Pedro Muñoz", "+56 9 5555 0001", date(2026, 10, 27),
     time(9, 30), time(7, 0), time(11, 0), "123456", "OC-4500123", 12, "No", None, None, "", "Andén 3"],
    ["Ferretería Los Aromos", "Pje. Los Aromos 455", "Puente Alto", "Luis Bravo", "+56 9 5555 0002", date(2026, 10, 27),
     time(11, 45), None, None, "123457", "", 4, "Sí", 22500, "Efectivo", "", ""],
]
libro(ejemplo, instrucciones=True).save("apps/web/public/plantilla-entregas.xlsx")

# --- Archivo de prueba: 50 filas. Resultado esperado: 41 listas, 5 por revisar y 4 con error
# (la guía G-50001 repetida marca sus DOS filas: no se sabe cuál es la buena). -----
cadenas = ["Jumbo", "Líder", "Santa Isabel", "Unimarc", "Tottus", "Easy", "Sodimac", "Acuenta"]
comunas = ["Providencia", "Ñuñoa", "La Florida", "Puente Alto", "Las Condes", "Maipú", "Santiago", "Macul"]
fecha = date(2030, 1, 15)
filas = []
for i in range(42):
    cadena = cadenas[i % len(cadenas)]
    comuna = comunas[(i // len(cadenas) + i) % len(comunas)]
    filas.append([f"{cadena} {comuna} {i + 1}", f"Av. Ficticia {100 + i * 37}", comuna, f"Recepción {i + 1}",
                  f"+56 9 6000 {1000 + i}", fecha, time(8 + (i % 4), 0), time(7, 0), time(13, 0), f"G-{50001 + i}",
                  f"OC-{9000 + i}", 1 + (i % 15), "No", None, None, "Matías Rojas" if i % 2 == 0 else "", ""])
# Por revisar (5)
filas.append(["Unimarc Viña", "Av. Libertad 100", "Viña del Mar", "Ana", "+56 9 6000 2001", fecha, time(9, 0), None, None, "G-60001", "", 3, "No", None, None, "", "Comuna fuera de la RM"])
filas.append(["Tottus Macul Norte", "Av. Macul 2000", "Macul", "", "", fecha, time(9, 0), None, None, "G-60002", "", 2, "No", None, None, "", "Sin teléfono"])
filas.append(["Easy Maipú Sur", "Av. Pajaritos 3000", "Maipú", "Rosa", "+56 9 6000 2003", fecha, time(9, 0), None, None, "G-60003", "", 5, "No", None, None, "Conductor Inexistente", ""])
filas.append(["Líder Santiago Centro", "Alameda 1500", "Santiago", "Juan", "+56 9 6000 2004", fecha, time(15, 0), time(7, 0), time(11, 0), "G-60004", "", 6, "No", None, None, "", "Hora fuera del horario"])
filas.append(["Acuenta Ñuñoa Sur", "Calle sin numero", "Ñuñoa", "Eva", "+56 9 6000 2005", fecha, time(10, 0), None, None, "G-60005", "", 1, "Sí", 45990, "Transferencia", "", "Dirección que no se encuentra"])
# Con error (3)
filas.append(["Jumbo Repetido", "Av. Ficticia 100", "Providencia", "X", "+56 9 6000 3001", fecha, time(9, 0), None, None, "G-50001", "", 2, "No", None, None, "", "Guía repetida"])
filas.append(["Santa Isabel Cero", "Av. Ficticia 999", "Macul", "X", "+56 9 6000 3002", fecha, time(9, 0), None, None, "G-70002", "", 0, "No", None, None, "", "Bultos en cero"])
filas.append(["", "Av. Ficticia 998", "Macul", "X", "+56 9 6000 3003", fecha, time(9, 0), None, None, "G-70003", "", 1, "No", None, None, "", "Sin cliente"])
assert len(filas) == 50
libro(filas).save("apps/web/e2e/archivos/entregas-50.xlsx")
print("listo")
