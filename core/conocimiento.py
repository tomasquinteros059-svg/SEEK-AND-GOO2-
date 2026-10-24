"""Base de conocimiento: notas, documentos del negocio y contactos.

Todo se guarda en SQLite. El texto de los documentos se extrae aquí mismo
(sin IA), y luego se entrega como contexto al escanear y al escribir mensajes.
"""
from __future__ import annotations

import csv
import io
import json
import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from core.validacion import normalizar_texto, parsear_fecha, parsear_monto

MAX_CARACTERES_ARCHIVO = 40_000
EXTENSIONES_ARCHIVO = {"pdf", "docx", "xlsx", "txt", "md", "csv"}


class ErrorConocimiento(Exception):
    pass


# ---------------- Notas ----------------

def listar_notas(conn: sqlite3.Connection) -> list[sqlite3.Row]:
    return conn.execute("SELECT * FROM notas ORDER BY id DESC").fetchall()


def agregar_nota(conn: sqlite3.Connection, texto: str, etiquetas: str = "") -> int:
    texto = texto.strip()
    if not texto:
        raise ErrorConocimiento("La nota está vacía.")
    cur = conn.execute("INSERT INTO notas (texto, etiquetas) VALUES (?, ?)", (texto, etiquetas.strip()))
    conn.commit()
    return cur.lastrowid


def actualizar_nota(conn: sqlite3.Connection, nota_id: int, texto: str, etiquetas: str | None = None) -> None:
    if not texto.strip():
        raise ErrorConocimiento("La nota no puede quedar vacía. Si no la necesitas, bórrala.")
    if etiquetas is None:
        conn.execute("UPDATE notas SET texto = ?, actualizado_en = datetime('now','localtime') WHERE id = ?",
                     (texto.strip(), nota_id))
    else:
        conn.execute("UPDATE notas SET texto = ?, etiquetas = ?, actualizado_en = datetime('now','localtime') WHERE id = ?",
                     (texto.strip(), etiquetas.strip(), nota_id))
    conn.commit()


def borrar_nota(conn: sqlite3.Connection, nota_id: int) -> None:
    conn.execute("DELETE FROM notas WHERE id = ?", (nota_id,))
    conn.commit()


# ---------------- Documentos del negocio ----------------

def _texto_celda(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        v = v.date()
    if isinstance(v, date):
        return v.strftime("%d-%m-%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _decodificar(datos: bytes) -> str:
    for codif in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return datos.decode(codif)
        except UnicodeDecodeError:
            continue
    return datos.decode("utf-8", errors="replace")


def extraer_texto(nombre: str, datos: bytes) -> str:
    """Saca el texto de un PDF, Word, Excel o archivo de texto."""
    ext = Path(nombre).suffix.lower().lstrip(".")
    if ext not in EXTENSIONES_ARCHIVO:
        raise ErrorConocimiento(f"«{nombre}»: formato no soportado. Usa PDF, Word (.docx), Excel (.xlsx) o TXT.")
    try:
        if ext in {"txt", "md", "csv"}:
            texto = _decodificar(datos)
        elif ext == "pdf":
            from pypdf import PdfReader

            lector = PdfReader(io.BytesIO(datos))
            texto = "\n\n".join((p.extract_text() or "") for p in lector.pages[:50])
        elif ext == "docx":
            from docx import Document

            doc = Document(io.BytesIO(datos))
            partes = [p.text for p in doc.paragraphs]
            for tabla in doc.tables:
                for fila in tabla.rows:
                    partes.append(" | ".join(c.text.strip() for c in fila.cells))
            texto = "\n".join(partes)
        else:  # xlsx
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
            partes = []
            for ws in wb.worksheets:
                partes.append(f"# {ws.title}")
                for fila in ws.iter_rows(values_only=True):
                    if any(v is not None for v in fila):
                        partes.append(" | ".join(_texto_celda(v) for v in fila))
            texto = "\n".join(partes)
    except ErrorConocimiento:
        raise
    except Exception as e:
        raise ErrorConocimiento(f"No se pudo leer «{nombre}». ¿Está dañado o protegido con clave?") from e
    texto = re.sub(r"\n{3,}", "\n\n", texto).strip()
    if not texto:
        raise ErrorConocimiento(f"«{nombre}» no tiene texto que se pueda leer. Si es un escaneo, súbelo en «Escanear».")
    return texto


def guardar_archivo(conn: sqlite3.Connection, nombre: str, datos: bytes) -> tuple[int, bool]:
    """Guarda el texto del archivo. Devuelve (id, recortado)."""
    texto = extraer_texto(nombre, datos)
    recortado = len(texto) > MAX_CARACTERES_ARCHIVO
    cur = conn.execute("INSERT INTO archivos_conocimiento (nombre, tipo, texto_extraido) VALUES (?, ?, ?)",
                       (nombre, Path(nombre).suffix.lower().lstrip("."), texto[:MAX_CARACTERES_ARCHIVO]))
    conn.commit()
    return cur.lastrowid, recortado


def listar_archivos(conn: sqlite3.Connection) -> list[sqlite3.Row]:
    return conn.execute("SELECT * FROM archivos_conocimiento ORDER BY id DESC").fetchall()


def borrar_archivo(conn: sqlite3.Connection, archivo_id: int) -> None:
    conn.execute("DELETE FROM archivos_conocimiento WHERE id = ?", (archivo_id,))
    conn.commit()


# ---------------- Contactos ----------------

CAMPOS_CONTACTO = {
    "nombre": ("Nombre", r"nombre|cliente|razon|contacto"),
    "telefono": ("Teléfono", r"tel|cel|fono|whats|movil"),
    "correo": ("Correo", r"correo|mail"),
    "deuda": ("Deuda", r"deuda|saldo|monto|debe|pendiente"),
    "fecha_vencimiento": ("Vencimiento", r"venc|fecha|pago"),
}


def normalizar_telefono(valor) -> str:
    """Deja los celulares chilenos como +56 9 1234 5678. Otros formatos quedan como venían."""
    if valor is None:
        return ""
    texto = _texto_celda(valor).strip()
    digitos = re.sub(r"\D", "", texto)
    if len(digitos) == 9 and digitos.startswith("9"):
        return f"+56 9 {digitos[1:5]} {digitos[5:]}"
    if len(digitos) == 8:
        return f"+56 9 {digitos[:4]} {digitos[4:]}"
    if len(digitos) == 11 and digitos.startswith("569"):
        return f"+56 9 {digitos[3:7]} {digitos[7:]}"
    return texto


def correo_valido(correo: str | None) -> bool:
    return bool(correo) and re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", correo.strip()) is not None


@dataclass
class Tabla:
    cabeceras: list[str]
    filas: list[list]


def leer_tabla(nombre: str, datos: bytes) -> Tabla:
    """Lee la primera hoja de un Excel, o un CSV (con coma o punto y coma)."""
    ext = Path(nombre).suffix.lower()
    try:
        if ext == ".xlsx":
            from openpyxl import load_workbook

            ws = load_workbook(io.BytesIO(datos), read_only=True, data_only=True).worksheets[0]
            filas = [list(f) for f in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in f)]
        elif ext in (".csv", ".txt"):
            texto = _decodificar(datos)
            primera = texto.splitlines()[0] if texto else ""
            sep = ";" if primera.count(";") > primera.count(",") else ","
            filas = [f for f in csv.reader(io.StringIO(texto), delimiter=sep) if any(c.strip() for c in f)]
        else:
            raise ErrorConocimiento("Usa un Excel (.xlsx) o un CSV.")
    except ErrorConocimiento:
        raise
    except Exception as e:
        raise ErrorConocimiento(f"No se pudo leer «{nombre}».") from e
    if len(filas) < 2:
        raise ErrorConocimiento("El archivo no tiene filas de datos (la primera fila debe ser de títulos).")
    ancho = max(len(f) for f in filas)
    cab = [(_texto_celda(c).strip() or f"Columna {i + 1}") for i, c in enumerate(filas[0] + [None] * (ancho - len(filas[0])))]
    return Tabla(cab, [f + [None] * (ancho - len(f)) for f in filas[1:]])


def adivinar_mapeo(cabeceras: list[str]) -> dict[str, int | None]:
    usados: set[int] = set()
    mapeo: dict[str, int | None] = {}
    for campo, (_, patron) in CAMPOS_CONTACTO.items():
        mapeo[campo] = None
        for i, cab in enumerate(cabeceras):
            if i not in usados and re.search(patron, normalizar_texto(cab) or cab.lower()):
                mapeo[campo] = i
                usados.add(i)
                break
    return mapeo


def contactos_desde_tabla(tabla: Tabla, mapeo: dict[str, int | None]) -> list[dict]:
    if mapeo.get("nombre") is None:
        raise ErrorConocimiento("Indica qué columna tiene el nombre.")
    usados = {i for i in mapeo.values() if i is not None}
    salida = []
    for fila in tabla.filas:
        def g(campo):
            i = mapeo.get(campo)
            return fila[i] if i is not None and i < len(fila) else None
        nombre = _texto_celda(g("nombre")).strip()
        if not nombre:
            continue
        venc = g("fecha_vencimiento")
        if isinstance(venc, datetime):
            venc = venc.date()
        fecha = venc if isinstance(venc, date) else (parsear_fecha(_texto_celda(venc)) if venc else None)
        extra = {tabla.cabeceras[i]: _texto_celda(v) for i, v in enumerate(fila) if i not in usados and v not in (None, "")}
        deuda = g("deuda")
        salida.append({
            "nombre": nombre,
            "telefono": normalizar_telefono(g("telefono")),
            "correo": _texto_celda(g("correo")).strip(),
            "deuda": parsear_monto(deuda if isinstance(deuda, (int, float)) else _texto_celda(deuda) or None),
            "fecha_vencimiento": fecha.isoformat() if fecha else None,
            "extra": extra,
        })
    return salida


def importar_contactos(conn: sqlite3.Connection, contactos: list[dict], *, reemplazar: bool = False,
                       origen: str = "") -> int:
    if reemplazar:
        conn.execute("DELETE FROM contactos")
    conn.executemany(
        "INSERT INTO contactos (nombre, telefono, correo, deuda, fecha_vencimiento, extra_json, origen) VALUES (?,?,?,?,?,?,?)",
        [(c["nombre"], c.get("telefono") or "", c.get("correo") or "", c.get("deuda"), c.get("fecha_vencimiento"),
          json.dumps(c.get("extra") or {}, ensure_ascii=False), origen) for c in contactos],
    )
    conn.commit()
    return len(contactos)


@dataclass
class Contacto:
    id: int
    nombre: str
    telefono: str
    correo: str
    deuda: int | None
    fecha_vencimiento: str | None
    extra: dict = field(default_factory=dict)

    @property
    def dias_atraso(self) -> int | None:
        f = parsear_fecha(self.fecha_vencimiento) if self.fecha_vencimiento else None
        return (date.today() - f).days if f else None


def listar_contactos(conn: sqlite3.Connection) -> list[Contacto]:
    return [Contacto(r["id"], r["nombre"], r["telefono"] or "", r["correo"] or "", r["deuda"], r["fecha_vencimiento"],
                     json.loads(r["extra_json"] or "{}"))
            for r in conn.execute("SELECT * FROM contactos ORDER BY nombre COLLATE NOCASE")]


def guardar_contacto(conn: sqlite3.Connection, c: dict, contacto_id: int | None = None) -> int:
    nombre = (c.get("nombre") or "").strip()
    if not nombre:
        raise ErrorConocimiento("Cada contacto necesita un nombre.")
    fecha = parsear_fecha(c.get("fecha_vencimiento")) if c.get("fecha_vencimiento") else None
    valores = (nombre, normalizar_telefono(c.get("telefono")), (c.get("correo") or "").strip(),
               parsear_monto(c.get("deuda")), fecha.isoformat() if fecha else None)
    if contacto_id is None:
        cur = conn.execute("INSERT INTO contactos (nombre, telefono, correo, deuda, fecha_vencimiento, origen) "
                           "VALUES (?,?,?,?,?, 'manual')", valores)
        contacto_id = cur.lastrowid
    else:
        conn.execute("UPDATE contactos SET nombre=?, telefono=?, correo=?, deuda=?, fecha_vencimiento=?, "
                     "actualizado_en=datetime('now','localtime') WHERE id=?", (*valores, contacto_id))
    conn.commit()
    return contacto_id


def borrar_contacto(conn: sqlite3.Connection, contacto_id: int) -> None:
    conn.execute("DELETE FROM contactos WHERE id = ?", (contacto_id,))
    conn.commit()


# ---------------- Contexto para la IA ----------------

def contexto_negocio(conn: sqlite3.Connection, max_caracteres: int = 12_000) -> str:
    """Notas y documentos del negocio, listos para incluir en las instrucciones de la IA."""
    partes = []
    notas = [f"- {n['texto']}" for n in reversed(listar_notas(conn))]
    if notas:
        partes.append("Notas del dueño:\n" + "\n".join(notas))
    for a in reversed(listar_archivos(conn)):
        partes.append(f"Documento «{a['nombre']}»:\n{a['texto_extraido'][:4000]}")
    texto = "\n\n".join(partes)
    if len(texto) > max_caracteres:
        texto = texto[:max_caracteres] + "\n[…]"
    return texto
