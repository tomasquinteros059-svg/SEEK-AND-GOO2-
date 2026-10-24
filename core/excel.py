"""Generación del Excel: una hoja por tipo de documento, detalle de productos y resumen.

Formato chileno: montos con signo $ y separador de miles (Excel en español de Chile
muestra $12.345), fechas dd-mm-aaaa, encabezados en negrita y filtros.
"""
from __future__ import annotations

import io
import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core import columnas, documentos
from core.validacion import limpiar_rut, normalizar_numero, normalizar_texto, parsear_fecha, parsear_monto

FMT_PESOS = '"$"#,##0'
FMT_FECHA = "dd-mm-yyyy"
HOJA_TRIB = columnas.TIPOS_DOCUMENTO["boleta_factura"]
HOJA_LISTAS = columnas.TIPOS_DOCUMENTO["lista"]
HOJA_FORMS = columnas.TIPOS_DOCUMENTO["formulario"]
HOJA_DETALLE = "Detalle"
HOJA_RESUMEN = "Resumen"
REF = "Ref."   # columna con el código interno del documento, para no duplicar al agregar
MONTOS = {"neto", "iva", "total", "precio_unitario"}

AZUL = PatternFill("solid", fgColor="2445A8")
CELESTE = PatternFill("solid", fgColor="E6EBF8")
BLANCO_NEGRITA = Font(bold=True, color="FFFFFF")


@dataclass
class Col:
    clave: str        # campo del documento, o "extra:<clave>"
    nombre: str


@dataclass
class ResultadoAgregar:
    contenido: bytes
    agregados: int
    omitidos: int


def ref(doc: dict) -> str:
    return f"DOC-{doc['id']}"


def _fecha(valor) -> date | str | None:
    f = parsear_fecha(valor) if valor else None
    return f or (valor or None)


def _columnas(conn: sqlite3.Connection, tipo: str) -> list[Col]:
    return [Col(c.clave if c.es_base else f"extra:{c.clave}", c.nombre_visible)
            for c in columnas.listar(conn, tipo, solo_activas=True)]


def _valor(doc: dict, clave: str):
    if clave.startswith("extra:"):
        return doc.get("extras", {}).get(clave[6:])
    if clave == "fecha":
        return _fecha(doc.get("fecha"))
    if clave == "tipo":
        return (doc.get("tipo") or "").capitalize()
    return doc.get(clave)


def _escribir_fila(ws: Worksheet, valores: list, claves: list[str]) -> None:
    ws.append(valores)
    fila = ws.max_row
    for i, clave in enumerate(claves, start=1):
        celda = ws.cell(fila, i)
        base = clave.split(":")[-1]
        if base in MONTOS and isinstance(celda.value, (int, float)):
            celda.number_format = FMT_PESOS
        elif isinstance(celda.value, (date, datetime)):
            celda.number_format = FMT_FECHA


def _dar_formato(ws: Worksheet) -> None:
    for celda in ws[1]:
        celda.font = BLANCO_NEGRITA
        celda.fill = AZUL
        celda.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if ws.max_column and ws.max_row:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 500)), start=1):
        largo = max((10 if isinstance(c.value, (date, datetime)) else len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(largo + 2, 9), 50)


def _hoja(wb: Workbook, nombre: str, cabeceras: list[str]) -> Worksheet:
    ws = wb.create_sheet(nombre)
    ws.append(cabeceras)
    return ws


def _llenar_libro(conn: sqlite3.Connection, wb: Workbook, docs: list[dict]) -> None:
    trib = [d for d in docs if d["tipo"] in ("boleta", "factura")]
    cols = _columnas(conn, "boleta_factura")
    ws = _hoja(wb, HOJA_TRIB, [c.nombre for c in cols] + [REF])
    for d in trib:
        _escribir_fila(ws, [_valor(d, c.clave) for c in cols] + [ref(d)], [c.clave for c in cols] + ["ref"])
    _dar_formato(ws)

    claves_det = ["fecha", "proveedor", "numero_documento", "descripcion", "cantidad", "precio_unitario", "total", "ref"]
    ws = _hoja(wb, HOJA_DETALLE, ["Fecha", "Proveedor", "N° documento", "Descripción", "Cantidad", "P. unitario", "Total", REF])
    for d in trib:
        for l in d.get("lineas", []):
            _escribir_fila(ws, [_fecha(d.get("fecha")), d.get("proveedor"), d.get("numero_documento"), l["descripcion"],
                                l["cantidad"], l["precio_unitario"], l["total"], ref(d)], claves_det)
    _dar_formato(ws)

    cols_l = _columnas(conn, "lista")
    ws = _hoja(wb, HOJA_LISTAS, [c.nombre for c in cols_l] + [REF])
    for d in (x for x in docs if x["tipo"] == "lista"):
        for it in d.get("items", []):
            fila = []
            for c in cols_l:
                if c.clave == "texto":
                    fila.append(it.get("texto"))
                elif c.clave == "cantidad":
                    fila.append(it.get("cantidad"))
                elif c.clave == "marcado":
                    fila.append("Sí" if it.get("marcado") else "No")
                else:
                    fila.append(_valor(d, c.clave))
            _escribir_fila(ws, fila + [ref(d)], [c.clave for c in cols_l] + ["ref"])
    _dar_formato(ws)

    forms = [d for d in docs if d["tipo"] == "formulario"]
    cols_f = _columnas(conn, "formulario")
    etiquetas: list[str] = []
    for f in forms:
        for c in f.get("campos", []):
            if c.get("etiqueta") and c["etiqueta"] not in etiquetas:
                etiquetas.append(c["etiqueta"])
    ws = _hoja(wb, HOJA_FORMS, [c.nombre for c in cols_f] + etiquetas + [REF])
    for f in forms:
        valores = {c.get("etiqueta"): c.get("valor") for c in f.get("campos", [])}
        _escribir_fila(ws, [_valor(f, c.clave) for c in cols_f] + [valores.get(e) for e in etiquetas] + [ref(f)],
                       [c.clave for c in cols_f] + ["texto"] * len(etiquetas) + ["ref"])
    _dar_formato(ws)


def _resumen(wb: Workbook, registros: list[dict]) -> None:
    """registros: dicts con fecha (date|None), proveedor, rut, neto, iva, total."""
    if HOJA_RESUMEN in wb.sheetnames:
        del wb[HOJA_RESUMEN]
    ws = wb.create_sheet(HOJA_RESUMEN)
    por_mes: dict[str, dict] = defaultdict(lambda: {"n": 0, "neto": 0, "iva": 0, "total": 0})
    por_prov: dict[str, dict] = defaultdict(lambda: {"rut": "", "n": 0, "total": 0})
    for r in registros:
        mes = r["fecha"].strftime("%Y-%m") if r.get("fecha") else "Sin fecha"
        m = por_mes[mes]
        m["n"] += 1
        for k in ("neto", "iva", "total"):
            m[k] += r.get(k) or 0
        p = por_prov[r.get("proveedor") or "Sin proveedor"]
        p["n"] += 1
        p["total"] += r.get("total") or 0
        p["rut"] = p["rut"] or (r.get("rut") or "")

    ws.append(["Totales por mes"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Mes", "Documentos", "Neto", "IVA", "Total"])
    cab1 = ws.max_row
    for mes in sorted(por_mes):
        m = por_mes[mes]
        ws.append([mes, m["n"], m["neto"], m["iva"], m["total"]])
        for c in "CDE":
            ws[f"{c}{ws.max_row}"].number_format = FMT_PESOS
    ws.append(["Total", len(registros), sum(r.get("neto") or 0 for r in registros),
               sum(r.get("iva") or 0 for r in registros), sum(r.get("total") or 0 for r in registros)])
    for c in "ABCDE":
        ws[f"{c}{ws.max_row}"].font = Font(bold=True)
    for c in "CDE":
        ws[f"{c}{ws.max_row}"].number_format = FMT_PESOS

    ws.append([])
    ws.append(["Totales por proveedor"])
    ws[f"A{ws.max_row}"].font = Font(bold=True, size=13)
    ws.append(["Proveedor", "RUT", "Documentos", "Total"])
    cab2 = ws.max_row
    for prov, p in sorted(por_prov.items(), key=lambda x: -x[1]["total"]):
        ws.append([prov, p["rut"], p["n"], p["total"]])
        ws[f"D{ws.max_row}"].number_format = FMT_PESOS
    for fila in (cab1, cab2):
        for celda in ws[fila]:
            if celda.value is not None:
                celda.font = Font(bold=True)
                celda.fill = CELESTE
    for letra, ancho in zip("ABCDE", (34, 14, 14, 14, 14)):
        ws.column_dimensions[letra].width = ancho


def _registros_desde_docs(docs: list[dict]) -> list[dict]:
    return [{"fecha": parsear_fecha(d.get("fecha")) if d.get("fecha") else None, "proveedor": d.get("proveedor"),
             "rut": d.get("rut_proveedor"), "neto": d.get("neto"), "iva": d.get("iva"), "total": d.get("total")}
            for d in docs if d["tipo"] in ("boleta", "factura")]


def _a_bytes(wb: Workbook) -> bytes:
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def crear_libro(conn: sqlite3.Connection) -> bytes:
    """Excel nuevo con todos los documentos guardados."""
    docs = documentos.obtener_todos(conn)
    wb = Workbook()
    wb.remove(wb.active)
    _llenar_libro(conn, wb, docs)
    _resumen(wb, _registros_desde_docs(docs))
    return _a_bytes(wb)


def _indice(ws: Worksheet) -> dict[str, int]:
    return {str(c.value).strip().lower(): c.column for c in ws[1] if c.value is not None}


def _valor_celda(ws: Worksheet, fila: int, col: int | None):
    return ws.cell(fila, col).value if col else None


def agregar_a_existente(conn: sqlite3.Connection, contenido: bytes) -> ResultadoAgregar:
    """Agrega a un Excel existente los documentos que aún no están en él."""
    try:
        wb = load_workbook(io.BytesIO(contenido))
    except Exception as e:  # archivo dañado o que no es .xlsx
        raise ValueError("No se pudo abrir ese archivo. ¿Es un Excel .xlsx válido?") from e
    docs = documentos.obtener_todos(conn)

    if HOJA_TRIB not in wb.sheetnames:
        # El archivo no tiene nuestras hojas: se agregan completas.
        _llenar_libro(conn, wb, docs)
        _resumen(wb, _registros_desde_docs(docs))
        return ResultadoAgregar(_a_bytes(wb), len(docs), 0)

    agregados = omitidos = 0
    ws = wb[HOJA_TRIB]
    idx = _indice(ws)
    nombres = {c.clave: c.nombre_visible.lower() for c in columnas.listar(conn, "boleta_factura")}
    c_ref, c_num = idx.get(REF.lower()), idx.get(nombres["numero_documento"])
    c_prov, c_rut = idx.get(nombres["proveedor"]), idx.get(nombres["rut_proveedor"])
    refs, claves = set(), set()
    for fila in range(2, ws.max_row + 1):
        if c_ref:
            refs.add(str(_valor_celda(ws, fila, c_ref) or ""))
        num = normalizar_numero(str(_valor_celda(ws, fila, c_num) or ""))
        if num:
            claves.add((num, "p", normalizar_texto(str(_valor_celda(ws, fila, c_prov) or ""))))
            claves.add((num, "r", limpiar_rut(str(_valor_celda(ws, fila, c_rut) or ""))))

    def ya_esta(d: dict) -> bool:
        if ref(d) in refs:
            return True
        num = normalizar_numero(d.get("numero_documento"))
        if not num:
            return False
        rut = limpiar_rut(d.get("rut_proveedor") or "")
        return (num, "p", normalizar_texto(d.get("proveedor"))) in claves or bool(rut and (num, "r", rut) in claves)

    trib_nuevos = []
    cols = _columnas(conn, "boleta_factura")
    for d in (x for x in docs if x["tipo"] in ("boleta", "factura")):
        if ya_esta(d):
            omitidos += 1
            continue
        trib_nuevos.append(d)
        # Se escribe en las columnas que ya tiene el archivo (respetando su orden y nombres)
        ancho = ws.max_column
        valores: list = [None] * ancho
        claves_fila = [""] * ancho
        for c in cols:
            pos = idx.get(c.nombre.lower())
            if pos:
                valores[pos - 1], claves_fila[pos - 1] = _valor(d, c.clave), c.clave
        if c_ref:
            valores[c_ref - 1] = ref(d)
        _escribir_fila(ws, valores, claves_fila)
        agregados += 1

    if HOJA_DETALLE in wb.sheetnames:
        det = wb[HOJA_DETALLE]
        claves_det = ["fecha", "proveedor", "numero_documento", "descripcion", "cantidad", "precio_unitario", "total", "ref"]
        for d in trib_nuevos:
            for l in d.get("lineas", []):
                _escribir_fila(det, [_fecha(d.get("fecha")), d.get("proveedor"), d.get("numero_documento"), l["descripcion"],
                                     l["cantidad"], l["precio_unitario"], l["total"], ref(d)], claves_det)

    # Listas y formularios: se agregan los que no están (por Ref.)
    otros = [d for d in docs if d["tipo"] in ("lista", "formulario")]
    for nombre_hoja, tipo in ((HOJA_LISTAS, "lista"), (HOJA_FORMS, "formulario")):
        if nombre_hoja not in wb.sheetnames:
            continue
        hoja = wb[nombre_hoja]
        col_ref = _indice(hoja).get(REF.lower())
        existentes = {str(hoja.cell(f, col_ref).value) for f in range(2, hoja.max_row + 1)} if col_ref else set()
        temporal = Workbook()
        temporal.remove(temporal.active)
        pendientes = [d for d in otros if d["tipo"] == tipo and ref(d) not in existentes]
        omitidos += sum(1 for d in otros if d["tipo"] == tipo and ref(d) in existentes)
        if not pendientes:
            continue
        _llenar_libro(conn, temporal, pendientes)
        origen = temporal[nombre_hoja]
        for fila in origen.iter_rows(min_row=2, values_only=True):
            hoja.append(list(fila))
            for celda in hoja[hoja.max_row]:
                if isinstance(celda.value, (date, datetime)):
                    celda.number_format = FMT_FECHA
        agregados += len(pendientes)

    # Resumen rehecho con TODAS las filas del archivo (las que ya tenía y las nuevas)
    idx = _indice(ws)
    registros = []
    for fila in range(2, ws.max_row + 1):
        def v(clave):
            return _valor_celda(ws, fila, idx.get(nombres[clave]))
        fecha = v("fecha")
        registros.append({
            "fecha": fecha.date() if isinstance(fecha, datetime) else (fecha if isinstance(fecha, date) else parsear_fecha(fecha)),
            "proveedor": v("proveedor"), "rut": v("rut_proveedor"),
            "neto": parsear_monto(v("neto")), "iva": parsear_monto(v("iva")), "total": parsear_monto(v("total")),
        })
    _resumen(wb, registros)
    return ResultadoAgregar(_a_bytes(wb), agregados, omitidos)


def actualizar_archivo_local(conn: sqlite3.Connection, ruta: Path) -> str | None:
    """Reescribe el Excel que se mantiene en la carpeta datos/. Devuelve un aviso si no se pudo."""
    try:
        ruta.write_bytes(crear_libro(conn))
        return None
    except PermissionError:
        return f"No se pudo actualizar {ruta.name}: ciérralo en Excel y vuelve a guardar."
    except OSError as e:
        return f"No se pudo actualizar {ruta.name}: {e}"
