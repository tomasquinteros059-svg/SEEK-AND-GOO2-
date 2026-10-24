import io
from datetime import date, timedelta

import pytest
from docx import Document
from openpyxl import Workbook

from core import conocimiento as k


def xlsx(filas) -> bytes:
    wb = Workbook(); ws = wb.active
    for f in filas:
        ws.append(f)
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def test_notas_crud(conn):
    i = k.agregar_nota(conn, "El cliente Juan Pérez paga siempre los días 5.")
    k.actualizar_nota(conn, i, "Juan Pérez paga los días 5 y 20.")
    assert k.listar_notas(conn)[0]["texto"] == "Juan Pérez paga los días 5 y 20."
    with pytest.raises(k.ErrorConocimiento):
        k.agregar_nota(conn, "   ")
    k.borrar_nota(conn, i)
    assert k.listar_notas(conn) == []


def test_extrae_texto_de_word_excel_y_txt(conn):
    d = Document(); d.add_paragraph("Pan amasado $2.500 la docena")
    t = d.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "Despacho"; t.rows[0].cells[1].text = "Gratis sobre $20.000"
    b = io.BytesIO(); d.save(b)
    assert "Gratis sobre $20.000" in k.extraer_texto("precios.docx", b.getvalue())
    assert "Hallulla | 1990" in k.extraer_texto("precios.xlsx", xlsx([["Producto", "Precio"], ["Hallulla", 1990]]))
    assert k.extraer_texto("nota.txt", "Horario: 7 a 20 h".encode("cp1252")) == "Horario: 7 a 20 h"
    with pytest.raises(k.ErrorConocimiento):
        k.extraer_texto("foto.jpg", b"x")
    with pytest.raises(k.ErrorConocimiento):
        k.extraer_texto("vacio.txt", b"   ")
    i, recortado = k.guardar_archivo(conn, "nota.txt", ("a" * 50_000).encode())
    assert recortado and len(k.listar_archivos(conn)[0]["texto_extraido"]) == k.MAX_CARACTERES_ARCHIVO


def test_extrae_texto_de_pdf():
    from pypdf import PdfWriter
    w = PdfWriter(); w.add_blank_page(200, 200); b = io.BytesIO(); w.write(b)
    with pytest.raises(k.ErrorConocimiento, match="no tiene texto"):
        k.extraer_texto("escaneo.pdf", b.getvalue())


def test_importa_contactos_desde_excel(conn):
    datos = xlsx([["Cliente", "Celular", "Email", "Saldo", "Fecha pago", "Comuna"],
                  ["Rosa Díaz", 987654321, "rosa@correo.cl", "$65.000", date(2026, 10, 30), "Ñuñoa"],
                  ["Luis Toro", None, "luis@correo.cl", 12000, "05-11-2026", "Maipú"],
                  [None, None, None, None, None, None]])
    tabla = k.leer_tabla("clientes.xlsx", datos)
    mapeo = k.adivinar_mapeo(tabla.cabeceras)
    assert mapeo == {"nombre": 0, "telefono": 1, "correo": 2, "deuda": 3, "fecha_vencimiento": 4}
    contactos = k.contactos_desde_tabla(tabla, mapeo)
    assert contactos[0] == {"nombre": "Rosa Díaz", "telefono": "+56 9 8765 4321", "correo": "rosa@correo.cl",
                            "deuda": 65000, "fecha_vencimiento": "2026-10-30", "extra": {"Comuna": "Ñuñoa"}}
    assert contactos[1]["telefono"] == "" and contactos[1]["fecha_vencimiento"] == "2026-11-05"
    assert k.importar_contactos(conn, contactos) == 2
    assert [c.nombre for c in k.listar_contactos(conn)] == ["Luis Toro", "Rosa Díaz"]
    k.importar_contactos(conn, contactos[:1], reemplazar=True)
    assert len(k.listar_contactos(conn)) == 1


def test_importa_csv_con_punto_y_coma():
    texto = "Nombre;Teléfono;Deuda\nAna Rojas;+56 9 1111 2222;72.500\n".encode("utf-8")
    tabla = k.leer_tabla("c.csv", texto)
    c = k.contactos_desde_tabla(tabla, k.adivinar_mapeo(tabla.cabeceras))[0]
    assert c["deuda"] == 72500 and c["telefono"] == "+56 9 1111 2222"


def test_editar_contacto_y_atraso(conn):
    i = k.guardar_contacto(conn, {"nombre": "Pedro", "telefono": "912345678", "deuda": "$10.000",
                                  "fecha_vencimiento": (date.today() - timedelta(days=4)).strftime("%d-%m-%Y")})
    c = k.listar_contactos(conn)[0]
    assert c.telefono == "+56 9 1234 5678" and c.deuda == 10000 and c.dias_atraso == 4
    k.guardar_contacto(conn, {"nombre": "Pedro Soto", "telefono": "", "deuda": None}, i)
    assert k.listar_contactos(conn)[0].nombre == "Pedro Soto"
    with pytest.raises(k.ErrorConocimiento):
        k.guardar_contacto(conn, {"nombre": ""})


def test_contexto_negocio(conn):
    k.agregar_nota(conn, "Juan paga los días 5.")
    k.guardar_archivo(conn, "precios.txt", "Hallulla $1.990 el kilo".encode())
    ctx = k.contexto_negocio(conn)
    assert "Juan paga los días 5." in ctx and "Hallulla $1.990" in ctx
    assert len(k.contexto_negocio(conn, max_caracteres=20)) <= 24
