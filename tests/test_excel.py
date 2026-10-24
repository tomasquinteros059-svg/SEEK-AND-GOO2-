import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from core import columnas, documentos, excel
from tests.conftest import respuesta


def cargar(conn):
    documentos.guardar_tributario(conn, respuesta("boleta_minimarket").documento_tributario, archivo_nombre="boleta.jpg")
    documentos.guardar_tributario(conn, respuesta("factura_distribuidora").documento_tributario, archivo_nombre="factura.jpg")
    documentos.guardar_lista(conn, respuesta("lista_feria_a_mano").lista)
    documentos.guardar_formulario(conn, respuesta("formulario_cliente").formulario)


def abrir(contenido: bytes):
    return load_workbook(io.BytesIO(contenido))


def filas(ws):
    return list(ws.iter_rows(values_only=True))


def test_hojas_y_formato(conn):
    cargar(conn)
    wb = abrir(excel.crear_libro(conn))
    assert wb.sheetnames == ["Boletas y Facturas", "Detalle", "Listas", "Formularios", "Resumen"]
    ws = wb["Boletas y Facturas"]
    assert [c.value for c in ws[1]] == ["Fecha", "Tipo", "Proveedor", "RUT", "N° documento", "Neto", "IVA", "Total", "Ref."]
    assert ws["A1"].font.b and ws.auto_filter.ref and ws.freeze_panes == "A2"
    assert ws["A2"].value == datetime(2026, 10, 14) and ws["A2"].number_format == "dd-mm-yyyy"
    assert ws["H2"].value == 16730 and ws["H2"].number_format == '"$"#,##0'
    assert ws["H3"].value == 256683 and ws["F3"].value == 215700
    assert len(filas(wb["Detalle"])) == 1 + 8
    assert len(filas(wb["Listas"])) == 1 + 6
    forms = filas(wb["Formularios"])
    assert "Teléfono" in forms[0] and "+56 9 8765 4321" in forms[1]


def test_resumen_por_mes_y_proveedor(conn):
    cargar(conn)
    resumen = filas(abrir(excel.crear_libro(conn))["Resumen"])
    assert ("2026-10", 2, 215700, 43654, 273413) in resumen
    assert ("Total", 2, 215700, 43654, 273413) in resumen
    assert ("DISTRIBUIDORA EL TRIGO LTDA.", "77.123.456-9", 1, 256683) in [r[:4] for r in resumen]


def test_columnas_configurables(conn):
    cargar(conn)
    cols = {c.clave: c for c in columnas.listar(conn, "boleta_factura")}
    columnas.actualizar(conn, cols["proveedor"].id, nombre_visible="Emisor")
    columnas.actualizar(conn, cols["tipo"].id, activa=False)
    nueva = columnas.agregar(conn, "boleta_factura", "Forma de pago", "efectivo o tarjeta")
    cab = [c.value for c in abrir(excel.crear_libro(conn))["Boletas y Facturas"][1]]
    assert "Emisor" in cab and "Proveedor" not in cab and "Tipo" not in cab and cab[-2] == "Forma de pago"
    with pytest.raises(ValueError):
        columnas.borrar(conn, cols["total"].id)
    columnas.borrar(conn, nueva.id)
    assert "Forma de pago" not in [c.value for c in abrir(excel.crear_libro(conn))["Boletas y Facturas"][1]]


def test_agregar_a_existente_sin_duplicar(conn):
    documentos.guardar_tributario(conn, respuesta("boleta_minimarket").documento_tributario)
    primero = excel.crear_libro(conn)
    documentos.guardar_tributario(conn, respuesta("factura_distribuidora").documento_tributario)
    documentos.guardar_lista(conn, respuesta("lista_feria_a_mano").lista)
    res = excel.agregar_a_existente(conn, primero)
    assert res.agregados == 2 and res.omitidos == 1
    wb = abrir(res.contenido)
    assert len(filas(wb["Boletas y Facturas"])) == 3
    assert len(filas(wb["Listas"])) == 1 + 6
    assert ("Total", 2, 215700, 43654, 273413) in filas(wb["Resumen"])
    # Volver a agregar no cambia nada
    otra = excel.agregar_a_existente(conn, res.contenido)
    assert otra.agregados == 0 and len(filas(abrir(otra.contenido)["Boletas y Facturas"])) == 3


def test_agregar_a_un_excel_ajeno(conn):
    cargar(conn)
    from openpyxl import Workbook
    wb = Workbook(); wb.active.title = "Mis cosas"; wb.active["A1"] = "no tocar"
    b = io.BytesIO(); wb.save(b)
    res = excel.agregar_a_existente(conn, b.getvalue())
    nuevo = abrir(res.contenido)
    assert nuevo["Mis cosas"]["A1"].value == "no tocar" and "Boletas y Facturas" in nuevo.sheetnames


def test_archivo_invalido(conn):
    with pytest.raises(ValueError):
        excel.agregar_a_existente(conn, b"esto no es un excel")
